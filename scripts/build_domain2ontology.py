#!/usr/bin/env python
"""Build the annotated domain→ontology release tables (``results/final/``).

The production matrix (``scripts/run_production_matrix.py``) leaves each
cell's significant associations keyed by opaque identifiers: InterPro
accessions in the domain column and raw vocabulary ids (GO:…, MP:…, MIM
numbers, CPX ids, …) in the term column. This script joins human-readable
names onto every cell and writes the release set:

    results/final/domain2-<cell>.tsv.gz    one annotated table per cell
    results/final/names/<table>_names.tsv  the id→name companion tables
    results/final/INDEX.md                 counts, coverage, provenance

The annotated tables are the input tables with two columns inserted —
``domain_name`` after ``domain`` (supra-domain names are the constituent
names joined with " + ") and ``term_name`` after the term column — every
other column byte-identical to the production output.

Every name comes from a file already on disk, in almost all cases the very
file the production run consumed (run manifests pin the hashes), so names and
ids belong to the same release:

* OBO ontologies (GO, ChEBI, DOID, MONDO, NCIT, HP, EFO, MP, WBPhenotype,
  WBbt, ZFA, FBcv, FBbt) — ``[Term]`` id/alt_id → name stanzas. EFO ids are
  normalised with the adapter's own ``normalise_efo_id``.
* EC — ``enzyme.dat`` for complete numbers plus ``enzclass.txt`` for the
  hierarchy interior ids ("3.1.-.-") minted by True Path propagation.
* SynGO — GO names where the term is a GO id; the SynGO release's own
  ``ontologies.xlsx`` for SYNGO:-private terms.
* Swiss-Prot flat file (one streaming pass) — OMIM disease names from
  ``CC -!- DISEASE`` blocks, Rhea reaction equations from ``CC -!- CATALYTIC
  ACTIVITY`` blocks, and ComplexPortal / DrugBank / TCDB / Reactome names
  from their ``DR`` lines.
* Orphanet — ORDO labels, topped up from ``en_product6.xml`` for codes newer
  than the ORDO release.
* Reactome — ``ReactomePathways.txt``, with the Swiss-Prot ``DR`` names
  covering pathways retired from the current Reactome release.
* OncoTree — the tumor-type JSON dump.
* Identity vocabularies (keyword, celltype, condensate, pharos) — the term id
  is already the curated display name; verified, not assumed.
* CAZy / MEROPS — the id *is* the standard nomenclature; class-level ids are
  decorated from small static maps documented inline.

Known, deliberate gaps (nothing on disk names these):

* MEROPS family/holotype ids (S01, T03.006, …) — needs a MEROPS download.
* UniPathway UPA ids — the resource is defunct; Swiss-Prot DR lines carry no
  labels.
* TCDB class/subclass ids ("1", "9.A") — DR lines carry family-level text
  only.
* GO:7770087 — a stray id in the all-species GAF, absent from go-basic.obo.

Unnamed ids get an empty name column; per-cell coverage is reported in
INDEX.md. Rerun is incremental: companion name tables are cached under
``results/final/names/`` (``--force`` rebuilds them).

Usage:
    uv run python scripts/build_domain2ontology.py
"""

from __future__ import annotations

import argparse
import gzip
import html
import io
import json
import re
import sys
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

from loguru import logger

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.gwas_annotation_source import normalise_efo_id  # noqa: E402

# ---------------------------------------------------------------------------
# Name-source inputs (the same files the production runs consumed).
# ---------------------------------------------------------------------------

NAME_INPUTS: dict[str, Path] = {
    "go_obo": Path("data/raw/go_ontology/go-basic.obo"),
    "chebi_obo": Path("data/raw/chebi/chebi_lite.obo"),
    "doid_obo": Path("data/raw/disease_ontology/doid.obo"),
    "mondo_obo": Path("data/raw/mondo/mondo.obo"),
    "ncit_obo": Path("data/raw/ncit/ncit.obo"),
    "hpo_obo": Path("data/raw/hpo/hp.obo"),
    "efo_obo": Path("data/raw/efo/efo.obo"),
    "mp_obo": Path("data/raw/mp_ontology/mp.obo"),
    "wbphenotype_obo": Path("data/raw/wormbase_ontology/wbphenotype.obo"),
    "wbbt_obo": Path("data/raw/wormbase_ontology/wbbt.obo"),
    "zfa_obo": Path("data/raw/zfin_ontology/zfa.obo"),
    "fbcv_obo": Path("data/raw/flybase_ontology/fbcv.obo"),
    "fbbt_obo": Path("data/raw/flybase_ontology/fbbt.obo"),
    "enzyme_dat": Path("data/raw/enzyme/enzyme.dat"),
    "enzclass": Path("data/raw/enzyme/enzclass.txt"),
    "subcell": Path("data/raw/uniprot_subcell/subcell.txt"),
    "syngo_zip": Path("data/raw/syngo/syngo1.3_complete_data.zip"),
    "uniprot_dat": Path("data/raw/uniprot_sprot_dat/uniprot_sprot.dat.gz"),
    "ordo_owl": Path("data/raw/orphanet/ordo_orphanet.owl"),
    "orphanet_xml": Path("data/raw/orphanet/en_product6.xml"),
    "reactome_pathways": Path("data/raw/reactome_relations/ReactomePathways.txt"),
    "oncotree_json": Path("data/raw/oncotree/oncotree_tumortypes.json"),
    "interpro_entries": Path("data/raw/interpro_entry.list"),
}

#: Vocabularies whose term id already is the curated display name.
IDENTITY_LABELS = frozenset({"keyword", "celltype", "condensate", "pharos"})

#: Nothing on disk names these (UniPathway is defunct upstream).
NAMELESS_LABELS = frozenset({"unipathway"})

#: MEROPS catalytic-type letters (alpha_prefix_ancestors interior nodes).
#: Standard MEROPS nomenclature; family/holotype ids stay unnamed — naming
#: them needs a MEROPS release file that is not on disk.
MEROPS_TYPES = {
    "A": "Aspartic peptidase",
    "C": "Cysteine peptidase",
    "G": "Glutamic peptidase",
    "I": "Peptidase inhibitor",
    "M": "Metallopeptidase",
    "N": "Asparagine peptide lyase",
    "P": "Peptidase of mixed catalytic type",
    "S": "Serine peptidase",
    "T": "Threonine peptidase",
    "U": "Peptidase of unknown catalytic type",
}

#: CAZy class prefixes (standard CAZy nomenclature; GT wording confirmed by
#: src/ontology_registry.py). Bare prefixes are alpha_prefix_ancestors
#: interior nodes from --propagate-annotations runs.
CAZY_CLASSES = {
    "GH": "Glycoside Hydrolase family",
    "GT": "GlycosylTransferase family",
    "PL": "Polysaccharide Lyase family",
    "CE": "Carbohydrate Esterase family",
    "AA": "Auxiliary Activity family",
    "CBM": "Carbohydrate-Binding Module family",
}

_CAZY_RE = re.compile(r"^([A-Za-z]+)(\d+)?$")


def cazy_name(term: str) -> str | None:
    """Decorate a CAZy family code with its class name; None if unknown."""
    match = _CAZY_RE.match(term)
    if match is None:
        return None
    class_name = CAZY_CLASSES.get(match.group(1))
    if class_name is None:
        return None
    return f"{class_name} {term}" if match.group(2) else class_name


# ---------------------------------------------------------------------------
# Parsers, one per source-file format.
# ---------------------------------------------------------------------------


def _join_wrapped(parts: list[str]) -> str:
    """Re-flow wrapped flat-file text, undoing mid-word hyphen wraps.

    UniProt CC blocks and Expasy DE lines wrap long tokens after a hyphen
    with no space; the continuation then starts with an alphanumeric, ``(``
    or ``[`` character and must rejoin without one.
    """
    text = ""
    for part in parts:
        if text.endswith("-") and part[:1] and (part[0].isalnum() or part[0] in "(["):
            text += part
        elif text:
            text += " " + part
        else:
            text = part
    return text


def parse_obo_names(
    path: Path,
    *,
    normalise: Callable[[str], str] | None = None,
    underscore_variants: bool = False,
) -> dict[str, str]:
    """``[Term]`` stanzas → {id: name}, alt_ids folded in, first-wins.

    ``normalise`` rewrites every id and alt_id (EFO stores its native ids
    idspace-prefixed, ``efo:EFO_0000001``). ``underscore_variants`` also emits
    ``PREFIX_LOCAL`` rows for non-numeric locals (``OBA:VT0000217`` →
    ``OBA_VT0000217``), which is how GWAS trait URIs spell them.
    """
    names: dict[str, str] = {}
    variants: dict[str, str] = {}
    in_term = False
    term_id: str | None = None
    term_name: str | None = None
    alt_ids: list[str] = []

    def flush() -> None:
        if term_id is None or term_name is None:
            return
        names.setdefault(term_id, term_name)
        for alt in alt_ids:
            variants.setdefault(alt, term_name)
        if underscore_variants and ":" in term_id:
            prefix, local = term_id.split(":", 1)
            if not local.isdigit():
                variants.setdefault(f"{prefix}_{local}", term_name)

    with open(path, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            line = line.rstrip("\n")
            if line.startswith("["):
                flush()
                in_term = line == "[Term]"
                term_id, term_name, alt_ids = None, None, []
            elif in_term and line.startswith("id: ") and term_id is None:
                term_id = normalise(line[4:].strip()) if normalise else line[4:].strip()
            elif in_term and line.startswith("name: ") and term_name is None:
                term_name = line[6:].strip()
            elif in_term and line.startswith("alt_id: "):
                alt = line[8:].strip()
                alt_ids.append(normalise(alt) if normalise else alt)
        flush()

    for key, value in variants.items():
        names.setdefault(key, value)
    return names


def parse_ec_names(enzyme_dat: Path, enzclass: Path) -> dict[str, str]:
    """Complete EC numbers from enzyme.dat, interior ids from enzclass.txt.

    enzyme.dat DE lines can wrap — mid-token after a hyphen, like the CC
    blocks, so the same hyphen-aware rejoin applies. Transferred/Deleted
    placeholder records are skipped. enzclass lines look like ``3. 1.-.-
    Acting on ester bonds.`` — the dotted number is de-spaced to match
    ``ec_ancestors`` interior ids.
    """
    names: dict[str, str] = {}
    ec_id: str | None = None
    description: list[str] = []

    def flush() -> None:
        if ec_id is None or not description:
            return
        text = _join_wrapped(description).rstrip(".")
        if not text.startswith(("Transferred entry", "Deleted entry")):
            names.setdefault(ec_id, text)

    with open(enzyme_dat, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith("ID   "):
                ec_id = line[5:].strip()
                description = []
            elif line.startswith("DE   "):
                description.append(line[5:].strip())
            elif line.startswith("//"):
                flush()
                ec_id, description = None, []
        flush()

    class_re = re.compile(r"^(\d+)\.\s*([\d-]+)\.\s*([\d-]+)\.\s*([\d-]+)\s+(.+)$")
    with open(enzclass, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            match = class_re.match(line.strip())
            if match:
                interior = ".".join(match.group(1, 2, 3, 4))
                names.setdefault(interior, match.group(5).strip().rstrip("."))
    return names


def parse_syngo_names(syngo_zip: Path, go_names: dict[str, str]) -> dict[str, str]:
    """SynGO term names: go-basic.obo where the id is GO, else ontologies.xlsx.

    The xlsx name column suffixes GO ids in parentheses ("synaptic vesicle
    (GO:0008021)"), so GO names are preferred for consistent spelling.
    """
    from openpyxl import load_workbook

    with zipfile.ZipFile(syngo_zip) as archive:
        workbook = load_workbook(
            io.BytesIO(archive.read("ontologies.xlsx")), read_only=True
        )
    names: dict[str, str] = {}
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip().lower() if cell else "" for cell in next(rows, ())]
        if "id" not in header or "name" not in header:
            continue
        id_col, name_col = header.index("id"), header.index("name")
        for row in rows:
            term_id = str(row[id_col]).strip() if row[id_col] is not None else ""
            name = str(row[name_col]).strip() if row[name_col] is not None else ""
            if term_id and name:
                names.setdefault(term_id, go_names.get(term_id, name))
        break
    workbook.close()
    return names


def parse_orphanet_names(ordo_owl: Path, product6_xml: Path) -> dict[str, str]:
    """ORPHA code → label: ORDO first, en_product6.xml for newer codes."""
    names: dict[str, str] = {}
    class_re = re.compile(
        r'<Class rdf:about="http://www\.orpha\.net/ORDO/Orphanet_(\d+)"'
    )
    label_re = re.compile(r"<rdfs:label(?:\s[^>]*)?>(.*?)</rdfs:label>")
    current: str | None = None
    with open(ordo_owl, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            match = class_re.search(line)
            if match:
                current = match.group(1)
                continue
            if "</Class>" in line:
                current = None
                continue
            if current is not None and current not in names:
                label = label_re.search(line)
                if label:
                    names[current] = html.unescape(label.group(1))

    code_re = re.compile(r"<OrphaCode>(\d+)</OrphaCode>")
    name_re = re.compile(r'<Name lang="en">(.*?)</Name>')
    pending: str | None = None
    with open(product6_xml, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            match = code_re.search(line)
            if match:
                pending = match.group(1)
                continue
            if pending is not None:
                match = name_re.search(line)
                if match:
                    names.setdefault(pending, html.unescape(match.group(1)))
                    pending = None
    return names


def parse_subcellular_names(subcell: Path) -> dict[str, str]:
    """subcell.txt: SL-nnnn → location/topology/orientation name.

    ID (location), IT (topology) and IO (orientation) entries share the SL-
    accession space, so all three prefixes must set the pending name that the
    following AC line keys.
    """
    names: dict[str, str] = {}
    pending: str | None = None
    with open(subcell, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith(("ID   ", "IT   ", "IO   ")):
                pending = line[5:].strip().rstrip(".")
            elif line.startswith("AC   ") and pending:
                names.setdefault(line[5:].strip(), pending)
                pending = None
    return names


def parse_reactome_names(pathways: Path) -> dict[str, str]:
    """ReactomePathways.txt: headerless ``stable_id  name  species``."""
    names: dict[str, str] = {}
    with open(pathways, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) >= 2 and fields[0]:
                names.setdefault(fields[0], fields[1])
    return names


def parse_oncotree_names(oncotree_json: Path) -> dict[str, str]:
    """OncoTree tumor-type dump: flat node list, code → name.

    The virtual root ``TISSUE`` has no node record (nodes cite it only as a
    parent), so True-Path-propagated output needs the explicit rendering.
    """
    nodes = json.loads(oncotree_json.read_text())
    names: dict[str, str] = {}
    for node in nodes:
        code, name = node.get("code"), node.get("name")
        if code and name:
            names.setdefault(code, name)
    names.setdefault("TISSUE", "Tissue (OncoTree root)")
    return names


def parse_interpro_names(entry_list: Path) -> dict[str, str]:
    """interpro_entry.list: ENTRY_AC → ENTRY_NAME (header skipped)."""
    names: dict[str, str] = {}
    with open(entry_list, encoding="utf-8", errors="replace") as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) >= 3 and fields[0].startswith("IPR"):
                names[fields[0]] = fields[2]
    return names


# ---------------------------------------------------------------------------
# The Swiss-Prot pass: six name tables from one stream over the flat file.
# ---------------------------------------------------------------------------

_MIM_RE = re.compile(r"^(.*?)\s*\[MIM:(\d+)\]")
_ABBREV_RE = re.compile(r"\s*\([A-Za-z0-9/-]+\)$")
_ISOFORM_TAG_RE = re.compile(r"\.\s*\[[A-Z][A-Z0-9]{5,9}-\d+\]$")
_RHEA_MASTER_RE = re.compile(r"Reaction=(.*?);\s*Xref=Rhea:(RHEA:\d+)")
_RHEA_DIRECTION_RE = re.compile(
    r"PhysiologicalDirection=([A-Za-z -]+?);\s*Xref=Rhea:(RHEA:\d+)"
)

SWISSPROT_TABLES = ("disease", "rhea", "complex", "drugbank", "tcdb", "reactome_dr")


def build_swissprot_tables(uniprot_dat: Path) -> dict[str, dict[str, str]]:
    """One streaming pass over uniprot_sprot.dat.gz → all six name tables."""
    tables: dict[str, dict[str, str]] = {name: {} for name in SWISSPROT_TABLES}
    tcdb_raw: dict[str, str] = {}
    cc_topic: str | None = None
    cc_parts: list[str] = []

    def flush_cc() -> None:
        nonlocal cc_topic, cc_parts
        if cc_topic == "DISEASE":
            text = _join_wrapped(cc_parts)
            if not text.startswith("Note="):
                match = _MIM_RE.match(text)
                if match:
                    name = _ABBREV_RE.sub("", match.group(1))
                    tables["disease"].setdefault(match.group(2), name)
        elif cc_topic == "CATALYTIC ACTIVITY":
            text = re.sub(r"\s+", " ", _join_wrapped(cc_parts))
            master = _RHEA_MASTER_RE.search(text)
            if master:
                equation = master.group(1).strip()
                tables["rhea"].setdefault(master.group(2), equation)
                for direction, rhea_id in _RHEA_DIRECTION_RE.findall(text):
                    tables["rhea"].setdefault(rhea_id, f"{equation} ({direction})")
        cc_topic, cc_parts = None, []

    def handle_dr(payload: str) -> None:
        fields = payload.split("; ", 2)
        if len(fields) < 3:
            return
        database, identifier, rest = fields[0], fields[1], fields[2].strip()
        if database == "ComplexPortal":
            name = _ISOFORM_TAG_RE.sub("", rest.rstrip(".")).rstrip(".")
            tables["complex"].setdefault(identifier, name)
        elif database == "DrugBank":
            tables["drugbank"].setdefault(identifier, rest.rstrip("."))
        elif database == "TCDB":
            tcdb_raw.setdefault(identifier, rest.rstrip("."))
        elif database == "Reactome":
            # Isoform-specific DR lines carry the same '. [ACC-n]' molecule
            # tag ComplexPortal lines do; strip it here too.
            name = _ISOFORM_TAG_RE.sub("", rest.rstrip(".")).rstrip(".")
            tables["reactome_dr"].setdefault(identifier, name)

    with gzip.open(uniprot_dat, "rt", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith("CC   "):
                payload = line[5:].rstrip("\n")
                if payload.startswith("-!- "):
                    flush_cc()
                    topic, _, first = payload[4:].partition(":")
                    if topic in ("DISEASE", "CATALYTIC ACTIVITY"):
                        cc_topic = topic
                        if first.strip():
                            cc_parts.append(first.strip())
                elif payload.startswith("---"):
                    flush_cc()
                elif cc_topic is not None:
                    cc_parts.append(payload.strip())
            else:
                flush_cc()
                if line.startswith("DR   "):
                    handle_dr(line[5:].rstrip("\n"))

    # TCDB DR descriptions are family-level (3-component) text, identical for
    # every id under a family, so prefixes of depth >= 3 inherit it exactly.
    tcdb = tables["tcdb"]
    for identifier, description in tcdb_raw.items():
        tcdb[identifier] = description
    for identifier, description in tcdb_raw.items():
        components = identifier.split(".")
        for depth in range(3, len(components)):
            tcdb.setdefault(".".join(components[:depth]), description)
    return tables


# ---------------------------------------------------------------------------
# Provider registry: vocabulary label → id-lookup callable, cached on disk.
# ---------------------------------------------------------------------------


@dataclass
class NameSources:
    """Lazily builds, caches, and serves the per-vocabulary name tables."""

    inputs: dict[str, Path]
    cache_dir: Path
    force: bool = False
    _tables: dict[str, dict[str, str]] = field(default_factory=dict)

    def _cached(
        self, table: str, build: Callable[[], dict[str, str]]
    ) -> dict[str, str]:
        if table in self._tables:
            return self._tables[table]
        cache = self.cache_dir / f"{table}_names.tsv"
        if cache.exists() and not self.force:
            names: dict[str, str] = {}
            with open(cache, encoding="utf-8") as handle:
                for line in handle:
                    identifier, _, name = line.rstrip("\n").partition("\t")
                    names[identifier] = name
        else:
            logger.info(f"Building name table: {table}")
            # Whitespace normalisation is release policy: raw sources carry
            # doubled/trailing spaces (and could carry tabs, which would break
            # the TSV), and the annotated tables must not depend on whether a
            # run built the table or read it back from this cache.
            names = {
                identifier: re.sub(r"\s+", " ", name).strip()
                for identifier, name in build().items()
            }
            self.cache_dir.mkdir(parents=True, exist_ok=True)
            with open(cache, "w", encoding="utf-8") as handle:
                for identifier in sorted(names):
                    handle.write(f"{identifier}\t{names[identifier]}\n")
        self._tables[table] = names
        return names

    def _obo(self, table: str, input_key: str, **kwargs: object) -> dict[str, str]:
        return self._cached(
            table,
            lambda: parse_obo_names(self.inputs[input_key], **kwargs),  # type: ignore[arg-type]
        )

    def _swissprot(self, table: str) -> dict[str, str]:
        if table not in self._tables:
            missing = [
                name
                for name in SWISSPROT_TABLES
                if self.force or not (self.cache_dir / f"{name}_names.tsv").exists()
            ]
            if missing:
                logger.info(
                    "Streaming Swiss-Prot flat file for name tables: "
                    + ", ".join(missing)
                )
                built = build_swissprot_tables(self.inputs["uniprot_dat"])
                for name in missing:
                    self._cached(name, lambda name=name: built[name])
            for name in SWISSPROT_TABLES:
                self._cached(name, lambda: {})  # cache hit; builder unused
        return self._tables[table]

    def provider(self, label: str) -> Callable[[str], str | None]:
        """Return the id → name lookup for one vocabulary label."""
        if label in IDENTITY_LABELS:
            return lambda term: term
        if label in NAMELESS_LABELS:
            return lambda term: None
        if label == "merops":
            return MEROPS_TYPES.get
        if label == "cazy":
            return cazy_name

        table = self._table_for(label)
        return table.get

    def _table_for(self, label: str) -> dict[str, str]:
        obo_labels = {
            "go": ("go", "go_obo"),
            "ligand": ("chebi", "chebi_obo"),
            "cofactor": ("chebi", "chebi_obo"),
            "doid": ("doid", "doid_obo"),
            "orphanet_doid": ("doid", "doid_obo"),
            "mondo": ("mondo", "mondo_obo"),
            "orphanet_mondo": ("mondo", "mondo_obo"),
            "ncit": ("ncit", "ncit_obo"),
            "hpo": ("hp", "hpo_obo"),
            "mp": ("mp", "mp_obo"),
            "wbphenotype": ("wbphenotype", "wbphenotype_obo"),
            "wbbt": ("wbbt", "wbbt_obo"),
            "zfa": ("zfa", "zfa_obo"),
            "fbcv": ("fbcv", "fbcv_obo"),
            "fbbt": ("fbbt", "fbbt_obo"),
        }
        if label in obo_labels:
            table, input_key = obo_labels[label]
            return self._obo(table, input_key)
        if label == "efo":
            return self._obo(
                "efo", "efo_obo", normalise=normalise_efo_id, underscore_variants=True
            )
        if label == "ec":
            return self._cached(
                "ec",
                lambda: parse_ec_names(
                    self.inputs["enzyme_dat"], self.inputs["enzclass"]
                ),
            )
        if label == "subcellular":
            return self._cached(
                "subcellular",
                lambda: parse_subcellular_names(self.inputs["subcell"]),
            )
        if label == "syngo":
            go_names = self._obo("go", "go_obo")
            return self._cached(
                "syngo", lambda: parse_syngo_names(self.inputs["syngo_zip"], go_names)
            )
        if label == "orphanet":
            return self._cached(
                "orphanet",
                lambda: parse_orphanet_names(
                    self.inputs["ordo_owl"], self.inputs["orphanet_xml"]
                ),
            )
        if label == "reactome":
            pathways = self._cached(
                "reactome",
                lambda: parse_reactome_names(self.inputs["reactome_pathways"]),
            )
            fallback = self._swissprot("reactome_dr")
            return {**fallback, **pathways}
        if label == "oncotree":
            return self._cached(
                "oncotree", lambda: parse_oncotree_names(self.inputs["oncotree_json"])
            )
        if label in ("disease", "rhea", "complex", "drugbank", "tcdb"):
            return self._swissprot(label)
        raise KeyError(f"no name source registered for ontology label {label!r}")


def named(provider: Callable[[str], str | None], term: str) -> str | None:
    """Look a term up, falling back to naming comma-joined composite parts."""
    name = provider(term)
    if name is not None:
        return name
    if "," in term:
        parts = [provider(part) for part in term.split(",")]
        if all(part is not None for part in parts):
            return " / ".join(parts)  # type: ignore[arg-type]
    return None


# ---------------------------------------------------------------------------
# Cell processing.
# ---------------------------------------------------------------------------


@dataclass
class CellResult:
    cell: str
    label: str
    species: str
    evidence: str
    config: str
    rows: int
    distinct_domains: int
    distinct_terms: int
    named_terms: int
    sha256: str
    output: Path


def iter_cells(production_dir: Path) -> Iterator[tuple[str, Path, dict]]:
    """Yield (cell_name, associations_path, manifest) for every matrix cell."""
    for cell_dir in sorted(production_dir.iterdir()):
        if not cell_dir.is_dir() or cell_dir.name == "surprise":
            continue
        manifests = sorted(cell_dir.glob("run_manifest_*.json"))
        if not manifests:
            logger.warning(f"Skipping {cell_dir.name}: no run manifest")
            continue
        manifest = json.loads(manifests[0].read_text())
        label = manifests[0].stem.removeprefix("run_manifest_")
        associations = cell_dir / f"domain_{label}_associations_significant.tsv"
        if not associations.exists():
            logger.warning(f"Skipping {cell_dir.name}: {associations.name} missing")
            continue
        yield cell_dir.name, associations, manifest


def annotate_cell(
    cell: str,
    associations: Path,
    manifest: dict,
    sources: NameSources,
    domain_names: dict[str, str],
    output_dir: Path,
) -> CellResult:
    """Join names onto one cell's associations and write the release table."""
    label = manifest["parameters"]["ontology"]
    provider = sources.provider(label)
    domains: set[str] = set()
    terms: dict[str, str | None] = {}
    rows = 0

    output = output_dir / f"domain2-{cell}.tsv.gz"
    with (
        open(associations, encoding="utf-8") as reader,
        gzip.GzipFile(output, "wb", mtime=0) as raw,
        io.TextIOWrapper(raw, encoding="utf-8") as writer,
    ):
        header = reader.readline().rstrip("\n").split("\t")
        if len(header) < 2:
            raise SystemExit(
                f"{associations} is empty or truncated — not a production "
                "association table; refusing to write a release file for it"
            )
        writer.write(
            "\t".join([header[0], "domain_name", header[1], "term_name", *header[2:]])
            + "\n"
        )
        for line in reader:
            if not line.strip():
                continue
            fields = line.rstrip("\n").split("\t")
            domain, term = fields[0], fields[1]
            domain_name = " + ".join(
                domain_names.get(part, part) for part in domain.split(",")
            )
            if term not in terms:
                terms[term] = named(provider, term)
            term_name = terms[term] or ""
            writer.write(
                "\t".join([domain, domain_name, term, term_name, *fields[2:]]) + "\n"
            )
            domains.add(domain)
            rows += 1

    sha256 = next(
        (
            out["sha256"]
            for out in manifest.get("outputs", [])
            if out.get("role") == "significant_associations"
        ),
        "",
    )
    params = manifest["parameters"]
    return CellResult(
        cell=cell,
        label=label,
        species=params.get("species", "human"),
        evidence=params.get("evidence_filter", "manual"),
        config="paper-parity" if params.get("propagate_annotations") else "baseline",
        rows=rows,
        distinct_domains=len(domains),
        distinct_terms=len(terms),
        named_terms=sum(1 for name in terms.values() if name is not None),
        sha256=sha256,
        output=output,
    )


def write_index(results: list[CellResult], output_dir: Path, commits: set[str]) -> None:
    commit_note = ", ".join(sorted(c for c in commits if c)) or "unknown"
    lines = [
        "# Domain→ontology mapping release",
        "",
        f"Annotated significant domain→term associations for {len(results)}",
        "cells of the production matrix (`results/production/`, git",
        f"`{commit_note}`). Built by `scripts/build_domain2ontology.py`; the",
        "`sha256` column pins each cell's source association file via its run",
        "manifest. Term/domain id→name companion tables are in `names/`.",
        "",
        "Regenerate with: `uv run python scripts/build_domain2ontology.py`",
        "",
        "Columns are the production output with `domain_name` and `term_name`",
        'inserted (supra-domain names join constituents with " + "). Unnamed',
        "ids keep an empty name column; `named terms` counts distinct terms.",
        "",
        "| file | ontology | species | evidence | config | rows | domains | terms | named terms | source sha256 |",
        "|---|---|---|---|---|---:|---:|---:|---:|---|",
    ]
    for result in results:
        lines.append(
            f"| {result.output.name} | {result.label} | {result.species} "
            f"| {result.evidence} | {result.config} | {result.rows:,} "
            f"| {result.distinct_domains:,} | {result.distinct_terms:,} "
            f"| {result.named_terms:,} | {result.sha256[:12]} |"
        )
    lines += [
        "",
        "## Known naming gaps",
        "",
        "Nothing on disk names these; ids are left with an empty name column:",
        "",
        "- **MEROPS** family/holotype ids (`S01`, `T03.006`) — only the",
        "  catalytic-type letters are named; needs a MEROPS release file.",
        "- **UniPathway** `UPA` ids — the resource is defunct; Swiss-Prot DR",
        "  lines carry no labels.",
        "- **TCDB** class/subclass ids (`1`, `9.A`) — Swiss-Prot DR text is",
        "  family-level only.",
        "- **GO:7770087** — stray id in the all-species GAF, absent from",
        "  go-basic.obo (1 association row).",
        "",
    ]
    (output_dir / "INDEX.md").write_text("\n".join(lines))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--production-dir",
        type=Path,
        default=Path("results/production"),
        help="Production matrix directory (default: results/production)",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("results/final"),
        help="Release output directory (default: results/final)",
    )
    parser.add_argument(
        "--cells",
        nargs="*",
        help="Only rebuild these cells (default: every cell)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Rebuild cached name tables in <output-dir>/names/",
    )
    args = parser.parse_args(argv)

    missing_inputs = [str(path) for path in NAME_INPUTS.values() if not path.exists()]
    if missing_inputs:
        raise SystemExit(
            "Missing name-source inputs (run from the repo root; see "
            "scripts/download_data.py): " + ", ".join(missing_inputs)
        )

    args.output_dir.mkdir(parents=True, exist_ok=True)
    sources = NameSources(
        inputs=NAME_INPUTS, cache_dir=args.output_dir / "names", force=args.force
    )
    domain_names = parse_interpro_names(NAME_INPUTS["interpro_entries"])
    logger.info(f"Loaded {len(domain_names):,} InterPro entry names")

    results: list[CellResult] = []
    commits: set[str] = set()
    for cell, associations, manifest in iter_cells(args.production_dir):
        if args.cells and cell not in args.cells:
            continue
        result = annotate_cell(
            cell, associations, manifest, sources, domain_names, args.output_dir
        )
        commits.add(manifest.get("git", {}).get("commit", ""))
        coverage = (
            f"{result.named_terms}/{result.distinct_terms}"
            if result.distinct_terms
            else "empty"
        )
        logger.info(
            f"{result.output.name}: {result.rows:,} rows, named terms {coverage}"
        )
        results.append(result)

    if args.cells:
        unmatched = set(args.cells) - {result.cell for result in results}
        if unmatched:
            raise SystemExit(
                f"No such production cell(s): {', '.join(sorted(unmatched))}"
            )
        logger.warning(
            "Partial rebuild (--cells): INDEX.md not regenerated — run without "
            "--cells to refresh it"
        )
    else:
        write_index(results, args.output_dir, commits)
        logger.info(f"Wrote {args.output_dir / 'INDEX.md'} ({len(results)} cells)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
